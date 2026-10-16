#!/usr/bin/env python3
"""Builds the docxtpl SOURCE template for the Supporting Evidence document —
the full record behind one Executive Summary Decision Request.

This produces a .docx with {{ jinja }} placeholders (not a filled
document), styled to match the RDR-CR-010 Supporting Evidence reference:
Azo Sans, teal headings, cream label-fill tables.

Run once to (re)generate app/templates/exports/risk_supporting_evidence_template.docx:

    python3 scripts/build_supporting_evidence_template.py

Edit THIS script (not the .docx by hand) when the template needs to
change, so the Jinja tags never get fragmented across Word runs. Paired
with scripts/build_exec_summary_template.py — both draw their context
from app/decision_request.py's build_context().
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _docx_style import (  # noqa: E402
    TEAL, GRAY, BODY_TXT, CREAM_FILL, FONT,
    add_table, add_data_row,
    label_value_table, h1, h2, body, body_multi, tag_p, spacer,
    add_header_footer, new_letter_doc,
)
from docx.shared import Pt

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_PATH = os.path.join(REPO_ROOT, "app", "templates", "exports", "risk_supporting_evidence_template.docx")

doc = new_letter_doc()
section = doc.sections[0]
add_header_footer(
    section,
    "SUPPORTING EVIDENCE  |  {{ risk_id }}",
    "Prepared by {{ requestor_name }}  •  {{ request_date }}  •  {{ request_id }}",
)

# ── Title block ──────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run("SUPPORTING EVIDENCE")
r.font.name = FONT
r.font.size = Pt(14)
r.font.bold = True
r.font.color.rgb = BODY_TXT

p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(6)
r2 = p2.add_run("Risk Decision Request {{ risk_id }} — {{ exec_subject }}")
r2.font.name = FONT
r2.font.size = Pt(10)
r2.font.bold = True
r2.font.color.rgb = TEAL

body(
    doc,
    "Generated from Risk Register entry {{ risk_id }}. This document provides the supporting evidence for "
    "the one-page Executive Summary Decision Request ({{ request_id }}). It contains the full risk mechanics, "
    "treatment options, vulnerability detail, and RACI assignments referenced in that summary.",
    italic=True, size=8.5, color=GRAY, after=10,
)

# ── 1. Request Metadata ───────────────────────────────────────────────────
h1(doc, "1. Request Metadata")
h2(doc, "1a. Big Ask")
body(doc, "{{ big_ask }}", after=8)

meta_w = [2600, 7768]
label_value_table(doc, meta_w, [
    ["Request ID", "{{ request_id }}"],
    ["Risk ID", "{{ risk_id }}"],
    ["Date", "{{ request_date }}"],
    ["Requestor", "{{ requestor_name }}, {{ requestor_role }}"],
    ["Executive owner requested", "{{ executive_owner_requested }}"],
    ["Risk owner", "{{ risk_owner }}"],
    ["Asset / system", "{{ asset_system }}"],
    ["Business process", "{{ business_process }}"],
    ["Related incident", "{{ related_incident_ref }}"],
    ["NIST CSF reference", "{{ nist_combined }}"],
])
spacer(doc, 8)

# ── 2. Source Risk Record ─────────────────────────────────────────────────
h1(doc, "2. Source Risk Record")
body(doc, "Pulled from the Risk Register — this request did not exist without a corresponding entry.", italic=True, size=8.5, color=GRAY, after=4)
label_value_table(doc, meta_w, [
    ["Risk category", "{{ risk_category }}"],
    ["Asset / system", "{{ asset_system }}"],
    ["Business process", "{{ business_process }}"],
    ["Risk owner", "{{ risk_owner }}"],
    ["Date identified", "{{ date_identified }}"],
    ["Last reviewed", "{{ last_reviewed }}"],
    ["Review frequency", "{{ review_frequency }}"],
    ["Board reporting flag", "{{ board_reporting_flag }}"],
    ["Related incident reference", "{{ related_incident_ref }}"],
])
spacer(doc, 8)

# ── 3. Risk Narrative and Business Impact ─────────────────────────────────
h1(doc, "3. Risk Narrative and Business Impact")
tag_p(doc, "{%p for para in risk_narrative_paragraphs %}")
body(doc, "{{ para }}", after=8, size=9)
tag_p(doc, "{%p endfor %}")

h2(doc, "Business Impact Summary")
impact_w = [2600, 7768]
label_value_table(doc, impact_w, [
    ["Operations", "{{ impact_operations }}"],
    ["Financial", "{{ impact_financial }}"],
    ["Customers / stakeholders", "{{ impact_customers }}"],
    ["Legal / regulatory / reputation", "{{ impact_legal }}"],
])

doc.add_page_break()

# ── 4. Risk Mechanics ──────────────────────────────────────────────────────
h1(doc, "4. Risk Mechanics (CISO Assessment)")
body(doc, "Scores and formulas below are computed by the Risk Register's scoring engine (app/scoring.py).", italic=True, size=8.5, color=GRAY, after=4)
mech_w = [2600, 7768]
label_value_table(doc, mech_w, [
    ["Asset / scope", "{{ asset_system }}"],
    ["Threat", "{{ threat_description }}"],
    ["Threat actor type", "{{ threat_actor_type }}"],
    ["Attack vector", "{{ attack_vector }}"],
    ["Threat intel source", "{{ threat_intel_source }}"],
    ["Vulnerability", "{{ vulnerability_detail }}"],
    ["Likelihood", "{{ likelihood_val }} — {{ likelihood_label }}"],
    ["Business impact", "{{ impact_val }} — {{ impact_label }}"],
    ["Inherent risk", "{{ inherent_score }} (maximum)"],
    ["Existing controls", "{{ existing_controls }}"],
    ["Control type", "{{ control_type }}"],
    ["Control effectiveness", "{{ control_effectiveness_val }} — {{ control_effectiveness_label }}"],
    ["Last control test", "{{ last_control_test_date }} — {{ control_test_result }}"],
    ["Residual risk (current)", "{{ residual_score }}"],
    ["RAG status", "{{ rag_status }}"],
    ["Exceeds risk appetite?", "{{ exceeds_appetite }}"],
])
spacer(doc, 8)

# ── 5. Options for Treatment ───────────────────────────────────────────────
h1(doc, "5. Options for Treatment")
tag_p(doc, "{%p if treatment_note %}")
body(doc, "{{ treatment_note }}", italic=True, size=8.5, after=6)
tag_p(doc, "{%p endif %}")

opt_w = [1600, 3268, 2500, 2968]
t_opt = add_table(doc, opt_w, ["Option", "Summary", "Pros", "Trade-offs / risks if chosen"])
add_data_row(t_opt, opt_w, ["Accept", "{{ opt_accept_summary }}", "{{ opt_accept_pros }}", "{{ opt_accept_tradeoffs }}"])
add_data_row(t_opt, opt_w, ["Avoid", "{{ opt_avoid_summary }}", "{{ opt_avoid_pros }}", "{{ opt_avoid_tradeoffs }}"], fill=CREAM_FILL)
add_data_row(t_opt, opt_w, ["Mitigate / Remediate", "{{ opt_mitigate_summary }}", "{{ opt_mitigate_pros }}", "{{ opt_mitigate_tradeoffs }}"])
add_data_row(t_opt, opt_w, ["Transfer", "{{ opt_transfer_summary }}", "{{ opt_transfer_pros }}", "{{ opt_transfer_tradeoffs }}"], fill=CREAM_FILL)

doc.add_page_break()

# ── 6. Decision Requirements and RACI ─────────────────────────────────────
h1(doc, "6. Decision Requirements and RACI")
h2(doc, "Decision Requirements")
dr_w = [3600, 6768]
label_value_table(doc, dr_w, [
    ["Select treatment direction", "{{ treatment_direction }} (CISO recommendation, from register)"],
    ["Set priority level", "{{ priority_level }}"],
    ["Decision needed by", "{{ decision_needed_by }}"],
    ["Authorize funding and resources", "{{ funding_requested }}"],
    ["Risk acceptance sign-off (if Accept)", "{{ acceptance_signoff }}"],
])
spacer(doc, 8)

h2(doc, "RACI")
raci_w = [1600, 3200, 5568]
t_raci = add_table(doc, raci_w, ["Role", "Named person / group", "Expected contribution"])
add_data_row(t_raci, raci_w, ["Responsible", "{{ raci_responsible }}", "Executes agreed treatment plan; reports status and evidence."])
add_data_row(t_raci, raci_w, ["Accountable", "{{ raci_accountable }}", "Owns the outcome and approves implementation and residual risk."], fill=CREAM_FILL)
add_data_row(t_raci, raci_w, ["Consulted", "{{ raci_consulted }}", "Provide expert input before and during execution."])
add_data_row(t_raci, raci_w, ["Informed", "{{ raci_informed }}", "Receive decision, impact, and status updates."], fill=CREAM_FILL)
spacer(doc, 8)

# ── 7. Action / Evidence Tracker ──────────────────────────────────────────
h1(doc, "7. Action / Evidence Tracker")
gd_w = [3800, 2200, 1400, 2968]
t_gd = add_table(doc, gd_w, ["Action / evidence", "Owner", "Due", "Status"])
add_data_row(t_gd, gd_w, ["{{ action1_desc }}", "{{ action1_owner }}", "{{ action1_due }}", "{{ action1_status }}"])
add_data_row(t_gd, gd_w, ["{{ action2_desc }}", "{{ action2_owner }}", "{{ action2_due }}", "{{ action2_status }}"], fill=CREAM_FILL)
add_data_row(t_gd, gd_w, ["{{ action3_desc }}", "{{ action3_owner }}", "{{ action3_due }}", "{{ action3_status }}"])

doc.add_page_break()

# ── 8. Vulnerability Evidence ──────────────────────────────────────────────
h1(doc, "8. Vulnerability Evidence")
body(doc, "{{ vuln_summary }}", italic=True, size=8.5, after=6)

vuln_w = [1900, 1300, 3868, 1600, 1300]
t_vuln = add_table(doc, vuln_w, ["Vulnerability ID", "CVSS severity", "Remediation", "Vulnerable product", "Days open"])
# Row-repeat loop: the {%tr for/endfor%} marker rows below must each be a
# full, standalone <w:tr> containing only that tag — docxtpl strips the
# marker row's own <w:tr> wrapper and leaves bare jinja control flow, so
# the *data* row between them (a real <w:tr>) is what actually repeats.
# See scripts/_docx_style.py module docstring / decision_request.py for
# why this can't be a plain paragraph before/after the table.
add_data_row(t_vuln, vuln_w, ["{%tr for v in vuln_findings %}", "", "", "", ""])
add_data_row(t_vuln, vuln_w, ["{{ v.cve_id }}", "{{ v.severity }}", "{{ v.remediation }}", "{{ v.vulnerable_product }}", "{{ v.days_open }}"])
add_data_row(t_vuln, vuln_w, ["{%tr endfor %}", "", "", "", ""])
spacer(doc, 8)

# ── 9. Assumptions, Dependencies, Review Trigger ──────────────────────────
h1(doc, "9. Assumptions, Dependencies, and Review Trigger")
body_multi(doc, [("Assumptions: ", True, BODY_TXT), ("{{ assumptions }}", False, BODY_TXT)], after=6)
body_multi(doc, [("Dependencies: ", True, BODY_TXT), ("{{ dependencies }}", False, BODY_TXT)], after=6)
body_multi(doc, [("Review trigger: ", True, BODY_TXT), ("{{ review_trigger }}", False, BODY_TXT)], after=8)

h2(doc, "9a. Evidence Links")
body_multi(doc, [("Evidence links: ", True, BODY_TXT), ("{{ evidence_links }}", False, BODY_TXT)], after=10)

doc.add_page_break()

# ── 10. Risk Owner Sign-Off ────────────────────────────────────────────────
h1(doc, "10. Risk Owner Sign-Off")
body(
    doc,
    "By signing below, the risk owner named in this register entry ({{ risk_owner }}) confirms the risk "
    "mechanics, existing controls, and residual risk score above accurately reflect the current state of "
    "this risk. If \"Accept\" is the treatment selected by the Executive Core Team, this signature also "
    "constitutes formal acceptance of the residual risk described in this document, subject to the "
    "conditions and review date recorded in the Executive Summary Decision Request.",
    after=14,
)
body(doc, "Risk Owner Signature: ____________________________________     Date: ____________", after=6)
body(doc, "Print Name: {{ risk_owner }}", after=0)

doc.save(OUTPUT_PATH)
print(f"template written to {OUTPUT_PATH}")
