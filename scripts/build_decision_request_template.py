#!/usr/bin/env python3
"""Builds the docxtpl SOURCE template for the Risk Decision Request export.

This produces a .docx with {{ jinja }} placeholders (not a filled document).
It is styled to match "Risk Decision Request Template - v3 (Updated)" —
navy/teal palette, same table shading, same header/footer.

Run once to (re)generate app/templates/exports/risk_decision_request_template.docx:

    python3 scripts/build_decision_request_template.py

Edit THIS script (not the .docx by hand) when the template needs to change,
so the Jinja tags never get fragmented across Word runs.
"""
import os
import docx
from docx import Document
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_PATH = os.path.join(REPO_ROOT, "app", "templates", "exports", "risk_decision_request_template.docx")

# ── Palette / type — matches original Template 03 ──────────────────────
NAVY = RGBColor(0x17, 0x36, 0x5D)
TEAL = RGBColor(0x0F, 0x6B, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_TXT = RGBColor(0x1F, 0x1F, 0x1F)
HEADER_GRAY = RGBColor(0x66, 0x66, 0x66)
ACCENT_BORDER = "4F81BD"
GRID_BORDER = "BFBFBF"
CALLOUT_FILL = "EAF2F8"
ZEBRA_FILL = "F5F8FA"
NAVY_HEX = "17365D"

FONT = "Calibri"
PAGE_W_TWIPS = 10368  # 12240 - 936*2


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color=GRID_BORDER, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Twips(w)
    # tblGrid
    tbl = table._tbl
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gridCol, w in zip(grid.findall(qn("w:gridCol")), widths):
            gridCol.set(qn("w:w"), str(w))


def set_cell_text(cell, text, bold=False, color=BODY_TXT, size=10, italic=False, align=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.text = ""
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run


def add_table(doc, widths, header_labels=None):
    cols = len(widths)
    rows = 1 if header_labels else 0
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    if header_labels:
        for i, label in enumerate(header_labels):
            c = table.rows[0].cells[i]
            set_cell_bg(c, NAVY_HEX)
            set_cell_borders(c)
            set_cell_text(c, label, bold=True, color=WHITE, size=10)
    set_col_widths(table, widths)
    return table


def add_data_row(table, widths, values, fill=None, italic=False):
    row = table.add_row()
    for i, val in enumerate(values):
        c = row.cells[i]
        set_cell_borders(c)
        if fill:
            set_cell_bg(c, fill)
        set_cell_text(c, val, italic=italic, size=10)
    set_col_widths(table, widths)
    return row


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = NAVY
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = NAVY
    return p


def body(doc, text, italic=False, bold=False, after=6, size=9.5, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.bold = bold
    r.font.color.rgb = BODY_TXT
    return p


def body_multi(doc, runs, after=6):
    """runs: list of (text, bold, color) tuples in one paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    for text, bold, color in runs:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(9.5)
        r.font.bold = bold
        r.font.color.rgb = color
    return p


def page_break(doc):
    doc.add_page_break()


def spacer(doc, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)


def callout_table(doc, header_text, body_text):
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    c0 = table.rows[0].cells[0]
    set_cell_bg(c0, NAVY_HEX)
    set_cell_borders(c0)
    set_cell_text(c0, header_text, bold=True, color=WHITE, size=10)
    c1 = table.rows[1].cells[0]
    set_cell_bg(c1, CALLOUT_FILL)
    set_cell_borders(c1)
    set_cell_text(c1, body_text, italic=True, size=10)
    set_col_widths(table, [PAGE_W_TWIPS])
    return table


def add_header_footer(section, header_text, footer_text):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(header_text)
    hr.font.name = FONT
    hr.font.size = Pt(8)
    hr.font.color.rgb = HEADER_GRAY

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(footer_text)
    fr.font.name = FONT
    fr.font.size = Pt(8)
    fr.font.color.rgb = HEADER_GRAY


# ── Build ────────────────────────────────────────────────────────────
doc = Document()

section = doc.sections[0]
section.page_width = Twips(12240)
section.page_height = Twips(15840)
section.top_margin = Twips(792)
section.bottom_margin = Twips(792)
section.left_margin = Twips(936)
section.right_margin = Twips(936)
section.header_distance = Twips(720)
section.footer_distance = Twips(720)

add_header_footer(
    section,
    "RISK DECISION REQUEST  |  {{ risk_id }}",
    "Prepared by {{ requestor_name }}  •  {{ request_date }}  •  {{ request_id }}",
)

# Default doc font
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(9.5)
style.font.color.rgb = BODY_TXT

# ── Page 1: Title + Source risk record + 1 Big Ask + Why it matters ──
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
pPr = p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "8")
bottom.set(qn("w:space"), "4")
bottom.set(qn("w:color"), ACCENT_BORDER)
pBdr.append(bottom)
pPr.append(pBdr)
r = p.add_run("RISK DECISION REQUEST")
r.font.name = FONT
r.font.size = Pt(26)
r.font.bold = True
r.font.color.rgb = NAVY

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.paragraph_format.space_after = Pt(16)
r2 = p2.add_run("Generated from Skerry entry {{ risk_id }} — for Executive Core Team decision")
r2.font.name = FONT
r2.font.size = Pt(12)
r2.font.italic = True
r2.font.color.rgb = TEAL

body(doc, "Request ID: {{ request_id }}   |   Date: {{ request_date }}   |   Requestor: {{ requestor_name }}, {{ requestor_role }}   |   Executive owner requested: {{ executive_owner_requested }}", after=10)

callout_table(doc, "1 BIG ASK  |  What decision is needed?", "{{ big_ask }}")
spacer(doc, 10)

h2(doc, "1. Source Risk Record")
body(doc, "Pulled from Skerry — this request did not exist without a corresponding entry.", italic=True, size=8.5)
t = add_table(doc, [3200, 7168], ["Field", "Value"])
rows = [
    ("Risk ID", "{{ risk_id }}"),
    ("Risk category", "{{ risk_category }}"),
    ("Asset / system", "{{ asset_system }}"),
    ("Business process", "{{ business_process }}"),
    ("Risk owner", "{{ risk_owner }}"),
    ("Date identified", "{{ date_identified }}"),
    ("Last reviewed", "{{ last_reviewed }}"),
    ("Review frequency", "{{ review_frequency }}"),
    ("Board reporting flag", "{{ board_reporting_flag }}"),
    ("Related incident reference", "{{ related_incident_ref }}"),
]
for i, (k, v) in enumerate(rows):
    add_data_row(t, [3200, 7168], [k, v], fill=ZEBRA_FILL if i % 2 else None)
spacer(doc, 10)

h2(doc, "2. Why It Matters")
body(doc, "{{ why_it_matters_lead }}")
t2 = add_table(doc, [2800, 7568], ["Business impact dimension", "Exposure / consequence"])
add_data_row(t2, [2800, 7568], ["Operations", "{{ impact_operations }}"])
add_data_row(t2, [2800, 7568], ["Financial", "{{ impact_financial }}"], fill=ZEBRA_FILL)
add_data_row(t2, [2800, 7568], ["Customers / stakeholders", "{{ impact_customers }}"])
add_data_row(t2, [2800, 7568], ["Legal / regulatory / reputation", "{{ impact_legal }}"], fill=ZEBRA_FILL)

page_break(doc)

# ── Page 2: Risk mechanics ────────────────────────────────────────────
h2(doc, "3. Risk Mechanics (CISO Assessment)")
body(doc, "Scores and formulas below are computed by Skerry's scoring engine (app/scoring.py).", italic=True, size=8.5)
t3 = add_table(doc, [2800, 7568], ["Element", "Assessment"])
mech_rows = [
    ("Asset / scope", "{{ asset_system }}"),
    ("Threat", "{{ threat_description }}"),
    ("Threat actor type", "{{ threat_actor_type }}"),
    ("Attack vector", "{{ attack_vector }}"),
    ("Threat intel source", "{{ threat_intel_source }}"),
    ("Vulnerability", "{{ vulnerability_detail }}"),
    ("Likelihood", "{{ likelihood_val }} — {{ likelihood_label }}"),
    ("Business impact", "{{ impact_val }} — {{ impact_label }}"),
]
for i, (k, v) in enumerate(mech_rows):
    add_data_row(t3, [2800, 7568], [k, v], fill=ZEBRA_FILL if i % 2 else None)
add_data_row(t3, [2800, 7568], ["Inherent risk", "{{ inherent_score }}"], fill=CALLOUT_FILL)
more_rows = [
    ("Existing controls", "{{ existing_controls }}"),
    ("Control type", "{{ control_type }}"),
    ("Control effectiveness", "{{ control_effectiveness_val }} — {{ control_effectiveness_label }}"),
    ("Last control test", "{{ last_control_test_date }} — {{ control_test_result }}"),
]
for i, (k, v) in enumerate(more_rows):
    add_data_row(t3, [2800, 7568], [k, v], fill=ZEBRA_FILL if i % 2 else None)
add_data_row(t3, [2800, 7568], ["Residual risk (current)", "{{ residual_score }}"], fill=CALLOUT_FILL)
add_data_row(t3, [2800, 7568], ["RAG status", "{{ rag_status }}"])
add_data_row(t3, [2800, 7568], ["Exceeds risk appetite?", "{{ exceeds_appetite }}"], fill=ZEBRA_FILL)

page_break(doc)

# ── Page 3: Options + Requested decision ──────────────────────────────
h2(doc, "4. Options for Treatment")
opt_w = [1400, 3000, 2800, 3168]
t4 = add_table(doc, opt_w, ["Option", "Summary", "Pros", "Trade-offs / risks if chosen"])
add_data_row(t4, opt_w, ["Accept", "{{ opt_accept_summary }}", "{{ opt_accept_pros }}", "{{ opt_accept_tradeoffs }}"])
add_data_row(t4, opt_w, ["Avoid", "{{ opt_avoid_summary }}", "{{ opt_avoid_pros }}", "{{ opt_avoid_tradeoffs }}"], fill=ZEBRA_FILL)
add_data_row(t4, opt_w, ["Mitigate / Remediate", "{{ opt_mitigate_summary }}", "{{ opt_mitigate_pros }}", "{{ opt_mitigate_tradeoffs }}"])
add_data_row(t4, opt_w, ["Transfer", "{{ opt_transfer_summary }}", "{{ opt_transfer_pros }}", "{{ opt_transfer_tradeoffs }}"], fill=ZEBRA_FILL)
spacer(doc, 8)

h2(doc, "5. Requested Decision and Priority")
t5 = add_table(doc, [3200, 7168], ["Decision needed from Executive Core Team", "Details"])
add_data_row(t5, [3200, 7168], ["Select treatment direction", "{{ treatment_direction }} (CISO recommendation, from register)"])
add_data_row(t5, [3200, 7168], ["Set priority level", "{{ priority_level }}"], fill=ZEBRA_FILL)
add_data_row(t5, [3200, 7168], ["Decision needed by", "{{ decision_needed_by }}"])
add_data_row(t5, [3200, 7168], ["Authorize funding and resources", "{{ funding_requested }}"], fill=ZEBRA_FILL)
add_data_row(t5, [3200, 7168], ["Risk acceptance sign-off (if Accept)", "{{ acceptance_signoff }}"])

page_break(doc)

# ── Page 4: RACI + Go deeper + refs ───────────────────────────────────
h2(doc, "6. RACI If Approved")
raci_w = [1600, 3200, 5568]
t6 = add_table(doc, raci_w, ["Role", "Named person / group", "Expected contribution"])
add_data_row(t6, raci_w, ["Responsible", "{{ raci_responsible }}", "Executes agreed treatment plan; reports status and evidence."])
add_data_row(t6, raci_w, ["Accountable", "{{ raci_accountable }}", "Owns the outcome and approves implementation and residual risk."], fill=ZEBRA_FILL)
add_data_row(t6, raci_w, ["Consulted", "{{ raci_consulted }}", "Provide expert input before and during execution."])
add_data_row(t6, raci_w, ["Informed", "{{ raci_informed }}", "Receive decision, impact, and status updates."], fill=ZEBRA_FILL)
spacer(doc, 8)

h2(doc, "7. Go Deeper (Supporting Detail)")
gd_w = [3800, 2200, 1400, 2968]
t7 = add_table(doc, gd_w, ["Action / evidence", "Owner", "Due", "Status"])
add_data_row(t7, gd_w, ["{{ action1_desc }}", "{{ action1_owner }}", "{{ action1_due }}", "{{ action1_status }}"])
add_data_row(t7, gd_w, ["{{ action2_desc }}", "{{ action2_owner }}", "{{ action2_due }}", "{{ action2_status }}"], fill=ZEBRA_FILL)
add_data_row(t7, gd_w, ["{{ action3_desc }}", "{{ action3_owner }}", "{{ action3_due }}", "{{ action3_status }}"])
spacer(doc, 10)

body_multi(doc, [("NIST CSF reference: ", True, BODY_TXT), ("Function: {{ nist_function }}   Category: {{ nist_category }}   Subcategory: {{ nist_subcategory }}", False, BODY_TXT)])
body_multi(doc, [("Evidence links: ", True, BODY_TXT), ("{{ evidence_links }}", False, BODY_TXT)])
body_multi(doc, [("Related incident: ", True, BODY_TXT), ("{{ related_incident_ref }}", False, BODY_TXT)])
body_multi(doc, [("Assumptions: ", True, BODY_TXT), ("{{ assumptions }}", False, BODY_TXT)])
body_multi(doc, [("Dependencies: ", True, BODY_TXT), ("{{ dependencies }}", False, BODY_TXT)])
body_multi(doc, [("Review trigger: ", True, BODY_TXT), ("{{ review_trigger }}", False, BODY_TXT)])

page_break(doc)

# ── Page 5: Executive Core Team Decision Record ───────────────────────
h1(doc, "Executive Core Team Decision Record")
body(doc, "To be completed by the Executive Core Team", italic=True, after=12)
body(doc, "Decision selected:   ☐ Accept   ☐ Avoid   ☐ Mitigate / Remediate   ☐ Transfer", after=12)
body(doc, "Accountable executive: ____________________________________", after=12)
body(doc, "Conditions, priority, resources, and review date:", after=4)
body(doc, "____________________________________________________________________", after=16)
spacer(doc, 20)

h2(doc, "Risk Owner Sign-Off")
body(doc, "By signing below, the risk owner named in this register entry ({{ risk_owner }}) confirms the risk mechanics, existing controls, and residual risk score above accurately reflect the current state of this risk. If \"Accept\" is the treatment selected by the Executive Core Team above, this signature also constitutes formal acceptance of the residual risk described in this document, subject to the conditions and review date recorded above.", after=16)
body(doc, "Risk Owner Signature: ____________________________________   Date: ______________", after=6)
body(doc, "Print Name: {{ risk_owner }}", after=0)

doc.save(OUTPUT_PATH)
print(f"template written to {OUTPUT_PATH}")
