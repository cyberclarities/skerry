#!/usr/bin/env python3
"""Builds the docxtpl SOURCE template for the Executive Summary Decision
Request — the one-page document that goes to the Executive Core Team.

This produces a .docx with {{ jinja }} placeholders (not a filled
document), styled to match the RDR-CR-010 Executive Summary reference:
Azo Sans, teal/red/green status colors, cream label-fill tables, a pink
decision-block header.

Run once to (re)generate app/templates/exports/risk_exec_summary_template.docx:

    python3 scripts/build_exec_summary_template.py

Edit THIS script (not the .docx by hand) when the template needs to
change, so the Jinja tags never get fragmented across Word runs. Paired
with scripts/build_supporting_evidence_template.py — both draw their
context from app/decision_request.py's build_context().
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _docx_style import (  # noqa: E402
    TEAL, RED, GREEN, GRAY, BODY_TXT, WHITE, CREAM_FILL, PINK_FILL, FONT, PAGE_W_TWIPS,
    set_cell_bg, set_cell_borders, set_cell_text, add_table, add_data_row,
    label_value_table, h1, h2, body, body_multi, bullet, tag_p, spacer,
    add_header_footer, new_letter_doc,
)
from docx.shared import Pt, Twips

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_PATH = os.path.join(REPO_ROOT, "app", "templates", "exports", "risk_exec_summary_template.docx")

doc = new_letter_doc()
section = doc.sections[0]
add_header_footer(
    section,
    "EXECUTIVE SUMMARY DECISION REQUEST  |  {{ risk_id }}",
    "{{ request_id }}  •  Full risk mechanics, vulnerability evidence, treatment options, and RACI are in the Supporting Evidence document.",
)

# ── Title block ──────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run("EXECUTIVE SUMMARY DECISION REQUEST")
r.font.name = FONT
r.font.size = Pt(14)
r.font.bold = True
r.font.color.rgb = BODY_TXT

p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(6)
r2 = p2.add_run("{{ risk_id }} — {{ exec_subject }}")
r2.font.name = FONT
r2.font.size = Pt(10)
r2.font.bold = True
r2.font.color.rgb = TEAL

body_multi(doc, [
    ("Request ID: ", True, BODY_TXT), ("{{ request_id }}     ", False, BODY_TXT),
    ("Date: ", True, BODY_TXT), ("{{ request_date }}     ", False, BODY_TXT),
    ("Requestor: ", True, BODY_TXT), ("{{ requestor_name }}, {{ requestor_role }}", False, BODY_TXT),
], after=2, size=8)
body_multi(doc, [
    ("Decision needed by: ", True, RED), ("{{ decision_needed_by }}     ", True, RED),
    ("Priority: ", True, BODY_TXT), ("{{ priority_level }}", False, BODY_TXT),
], after=8, size=8)

# ── Executive Ask ────────────────────────────────────────────────────────
h2(doc, "Executive Ask")
body(doc, "{{ executive_ask }}", after=4, size=9)
body_multi(doc, [
    ("Recommended direction: ", True, BODY_TXT), ("{{ suggested_treatment }}", True, GREEN),
], after=8, size=9)

# ── Risk at a Glance ─────────────────────────────────────────────────────
h2(doc, "Risk at a Glance")
glance_w = [2400, 2784, 2400, 2784]
label_value_table(doc, glance_w, [
    ["Asset / System", "{{ asset_system }}", "Risk Owner", "{{ risk_owner }}"],
    ["Residual Risk", "{{ residual_score }} — {{ rag_status }}", "Inherent Risk", "{{ inherent_score }} (max)"],
    ["Likelihood", "{{ likelihood_val }} — {{ likelihood_label }}", "Business Impact", "{{ impact_val }} — {{ impact_label }}"],
    ["Exceeds Appetite", "{{ exceeds_appetite }}", "Financial Exposure", "{{ financial_exposure_abbrev }}"],
    ["Vulnerabilities", "{{ vuln_open_count }} open ({{ vuln_high_count }} High)", "Oldest Open", "{{ vuln_oldest_days }} days"],
])
spacer(doc, 4)

# ── Why This Matters ─────────────────────────────────────────────────────
h2(doc, "Why This Matters")
tag_p(doc, "{%p for b in why_it_matters_bullets %}")
bullet(doc, "{{ b }}")
tag_p(doc, "{%p endfor %}")
spacer(doc, 3)

# ── Recommended Action ───────────────────────────────────────────────────
h2(doc, "Recommended Action")
tag_p(doc, "{%p for a in recommended_action_bullets %}")
bullet(doc, "{{ a }}")
tag_p(doc, "{%p endfor %}")
spacer(doc, 4)

# ── Risk Decision block ──────────────────────────────────────────────────
dec = doc.add_table(rows=1, cols=1)
dec.style = "Table Grid"
dec.autofit = False
dec.columns[0].width = Twips(PAGE_W_TWIPS)
c0 = dec.rows[0].cells[0]
c0.width = Twips(PAGE_W_TWIPS)
set_cell_bg(c0, PINK_FILL)
set_cell_borders(c0)
set_cell_text(c0, "RISK DECISION — Executive Core Team", bold=True, color=BODY_TXT, size=10)


def decision_row(label_tpl, label_suffix_tpl, text_tpl):
    row = dec.add_row()
    c = row.cells[0]
    set_cell_borders(c)
    c.width = Twips(PAGE_W_TWIPS)
    p = c.paragraphs[0]
    p.text = ""
    r1 = p.add_run(f"☐  {label_tpl}")
    r1.font.name = FONT
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    # label_suffix_tpl renders as a docxtpl RichText (green ★ RECOMMENDED or
    # gray (not viable)) — RichText values require the {{r ...}} tag, not
    # {{ ... }}, or docxtpl silently drops them. The placeholder run's own
    # styling doesn't matter; RichText supplies its own runs on render.
    r2 = p.add_run(label_suffix_tpl)
    r2.font.name = FONT
    r2.font.size = Pt(9.5)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_before = Pt(1)
    r3 = p2.add_run(text_tpl)
    r3.font.name = FONT
    r3.font.size = Pt(9)
    return row


decision_row("Mitigate / Remediate", "{{r exec_opt_mitigate_label }}", "{{ exec_opt_mitigate_text }}")
decision_row("Accept", "{{r exec_opt_accept_label }}", "{{ exec_opt_accept_text }}")
decision_row("Avoid", "{{r exec_opt_avoid_label }}", "{{ exec_opt_avoid_text }}")
decision_row("Transfer", "{{r exec_opt_transfer_label }}", "{{ exec_opt_transfer_text }}")

cond_row = dec.add_row()
cc = cond_row.cells[0]
set_cell_borders(cc)
cc.width = Twips(PAGE_W_TWIPS)
set_cell_text(cc, "Decision conditions / required actions: __________________________________________________", size=9)

exec_row = dec.add_row()
ec = exec_row.cells[0]
set_cell_borders(ec)
ec.width = Twips(PAGE_W_TWIPS)
set_cell_text(ec, "Accountable executive: ____________________________     Review date: ________________", size=9)

spacer(doc, 6)

# ── Signatures ────────────────────────────────────────────────────────────
body(doc, "Executive approver signature: ____________________________________     Date: ____________", after=6, size=9)
body(doc, "Risk owner acknowledgement ({{ risk_owner }}): ____________________________________     Date: ____________", after=6, size=9)

body(
    doc,
    "Full risk mechanics, vulnerability evidence, treatment options, and RACI are in the Supporting "
    "Evidence document ({{ request_id }} Supporting Evidence).",
    italic=True, size=7.5, color=GRAY, after=0,
)

doc.save(OUTPUT_PATH)
print(f"template written to {OUTPUT_PATH}")
